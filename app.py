import streamlit as st
import pandas as pd
import plotly.express as px
import easyocr
import numpy as np
import cv2
import re
from PIL import Image
from io import BytesIO
from openpyxl import Workbook
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.utils.dataframe import dataframe_to_rows

# --- 1. 網頁初始配置 ---
st.set_page_config(page_title="汽車保養工單系統", layout="centered")

# 初始化 EasyOCR 辨識器 (繁體中文與英文)
@st.cache_resource
def load_ocr():
    return easyocr.Reader(['ch_tra', 'en'], gpu=False)

reader = load_ocr()

st.title("🚗 汽車保養工單系統 (多欄位 AI 擷取版)")
st.write("上傳工單照片，AI 將自動擷取：車牌號碼、總金額、行駛里程、保養項目明細。")

# --- 2. 初始化資料庫 (在記憶體中模擬) ---
if "order_list" not in st.session_state:
    st.session_state.order_list = [
        {"工單單號": "D260513-375", "車牌號碼": "ANV-1055", "行駛里程": "140,029", "保養項目": "215/55R17 Michelin PRIMACY 5, 輪胎拆裝工資, 輪胎平衡校正, 氮氣填充, 四輪定位-電腦3D, 煞車來令片(前/WTC), 力魔LM3318-快速清潔噴劑, 底盤拆裝工資, 煞車來令片(後/WTC), 煞車盤(後)", "保養類別": "輪胎定位", "金額": 24300},
    ]

# --- 3. 全域變數初始化 ---
extracted_wo = ""
extracted_plate = ""
extracted_mileage = ""
extracted_price = 0
extracted_items_str = "" 

# --- 4. AI 影像自動擷取區塊 ---
st.header("📸 AI 工單自動擷取")

input_method = st.radio("請選擇輸入方式：", ["從手機相簿上傳照片/檔案", "使用手機相機現場拍照"])
img_file = st.file_uploader("請選擇工單照片", type=["jpg", "jpeg", "png"]) if input_method == "從手機相簿上傳照片/檔案" else st.camera_input("請對準工單拍照")

if img_file is not None:
    image = Image.open(img_file)
    img_np = np.array(image)
    
    # 影像預處理：強化對比度
    gray = cv2.cvtColor(img_np, cv2.COLOR_RGB2GRAY)
    enhanced_img = cv2.equalizeHist(gray)
    
    with st.spinner("AI 正在分析工單資訊中..."):
        results = reader.readtext(enhanced_img, detail=1, paragraph=False)
        
        all_texts = []
        raw_lines = []
        for res in results:
            text_curr = res[1].strip().replace(" ", "")
            raw_lines.append(text_curr)
            all_texts.append(text_curr.upper())
            
        full_text_block = "||".join(all_texts)
        
        # 顯示除錯紀錄
        with st.expander("🔍 檢視 AI 影像除錯紀錄"):
            st.write(raw_lines)

        # ---- 核心關鍵字擷取演算法 ----
        
        # 1. 擷取車牌號碼
        plate_pattern = r'[A-Z0-9]{2,4}[-─][A-Z0-9]{2,4}'
        plates_found = re.findall(plate_pattern, full_text_block)
        if plates_found:
            extracted_plate = plates_found[0]
        else:
            for i, text in enumerate(all_texts):
                if any(k in text for k in ["車牌", "車號", "車輛"]):
                    for offset in range(0, 3):
                        if i + offset < len(all_texts):
                            candidate = re.sub(r'[^\w-]', '', all_texts[i + offset])
                            if len(candidate) >= 6 and any(c.isdigit() for c in candidate):
                                extracted_plate = candidate
                                break
                if extracted_plate: break

        # 2. 擷取行駛里程
        for i, text in enumerate(all_texts):
            if "里程" in text or "公里" in text:
                for offset in range(0, 3):
                    if i + offset < len(all_texts):
                        candidate = all_texts[i + offset]
                        mileage_match = re.search(r'[\d,]+', candidate)
                        if mileage_match and len(mileage_match.group(0)) >= 3:
                            extracted_mileage = mileage_match.group(0)
                            break
            if extracted_mileage: break

        # 3. 擷取總金額
        for i, text in enumerate(all_texts):
            if any(k in text for k in ["總金額", "金額", "總計", "合計"]):
                combined_look = "".join(all_texts[i:i+4])
                digits = re.findall(r'[\d,.]+', combined_look)
                for d_str in digits:
                    clean_d = d_str.replace(",", "").split(".")[0]
                    if clean_d.isdigit() and 100 < int(clean_d) < 1000000:
                        extracted_price = int(clean_d)
                        break
            if extracted_price > 0: break

        # 4. 【演算法大幅升級】：改用排除法精準擷取保養項目
        extracted_items = []
        # 設定要過濾掉的非保養項目雜訊關鍵字
        ignore_keywords = [
            "總金額", "金額", "總計", "合計", "車牌", "車號", "里程", "公里", 
            "馳加", "工單", "單號", "日期", "客戶", "電話", "地址", "統一編號", 
            "小計", "營業稅", "銷售額", "外銷", "新台幣", "備註", "應收"
        ]
        
        for text in raw_lines:
            # 條件 A：不能包含任何系統與金額的雜訊關鍵字
            if any(ik in text for ik in ignore_keywords):
                continue
                
            # 條件 B：排除純數字、純英文字（通常是料號、電話或統編）
            clean_text = re.sub(r'^[A-Z0-9\-_]{4,20}', '', text) # 去除行首的長串商品編號
            clean_text = re.sub(r'[^\w\u4e00-\u9fa5\/\-]', '', clean_text) # 只留中英文字、斜線、橫線
            
            # 條件 C：必須包含中文（保養品項必定有中文描述，如：輪胎、更換、工資）
            has_chinese = any('\u4e00' <= char <= '\u9fa5' for char in clean_text)
            
            if has_chinese and len(clean_text) >= 2:
                if clean_text not in extracted_items:
                    extracted_items.append(clean_text)
                    
        extracted_items_str = ", ".join(extracted_items)

        # 抓工單單號
        wo_found = re.search(r'D[0-9]{5,6}[-─][0-9]+', full_text_block)
        if wo_found: extracted_wo = wo_found.group(0)

        # ---- 畫面上秀出 AI 擷取結果 ----
        st.subheader("🎯 AI 欄位擷取結果")
        col_res1, col_res2, col_res3 = st.columns(3)
        col_res1.metric("車牌號碼", extracted_plate if extracted_plate else "未偵測到")
        col_res2.metric("行駛里程", f"{extracted_mileage} KM" if extracted_mileage else "未偵測到")
        col_res3.metric("辨識總金額", f"${extracted_price:,}" if extracted_price else "未偵測到")
        st.text_area("自動擷取保養項目清單確認", value=extracted_items_str, height=100)

# --- 5. 確認與填寫工單資料表單 ---
st.header("📝 確認與填寫工單資料")
with st.form("order_form", clear_on_submit=False):
    col1, col2 = st.columns(2)
    with col1:
        wo_number = st.text_input("工單單號", value=extracted_wo if extracted_wo else f"WO-2026{len(st.session_state.order_list)+1:03d}")
        car_number = st.text_input("車牌號碼", value=extracted_plate if extracted_plate else "", placeholder="例如: ANV-1055")
        mileage = st.text_input("行駛里程 (KM)", value=extracted_mileage if extracted_mileage else "", placeholder="例如: 140,029")
    with col2:
        category = st.selectbox("保養類別", ["輪胎定位", "底盤系統", "定期保養", "引擎系統", "冷氣系統", "其他"])
        price = st.number_input("金額 (元)", min_value=0, value=int(extracted_price), step=100)
        
    item_name = st.text_area("保養項目明細描述", value=extracted_items_str, placeholder="項目會由 AI 自動帶入，也可手動修改")
    
    # 表單送出按鈕
    submit_btn = st.form_submit_button("確認新增此筆工單至系統")
    
    if submit_btn:
        if car_number and price > 0:
            new_order = {
                "工單單號": wo_number,
                "車牌號碼": car_number,
                "行駛里程": mileage if mileage else "未記錄",
                "保養項目": item_name if item_name else "未填寫明細",
                "保養類別": category,
                "金額": price
            }
            st.session_state.order_list.append(new_order)
            st.success(f"🎉 工單 {wo_number} 已成功寫入系統資料庫！")
            st.rerun()
        else:
            st.error("❌ 請確認車牌號碼與金額是否正確填寫！")

# --- 6. 圖表與 Excel 匯出邏輯 ---
df_detail = pd.DataFrame(st.session_state.order_list)

st.header("📊 汽車保養類別金額統計圓餅圖")
if not df_detail.empty:
    df_summary = df_detail.groupby("保養類別")["金額"].sum().reset_index()
    fig = px.pie(df_summary, values='金額', names='保養類別', title='各類別累積消費比例', hole=0.3)
    fig.update_traces(textinfo='value+percent')
    st.plotly_chart(fig, use_container_width=True)

def generate_excel(df_d, df_s):
    output = BytesIO()
    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "類別統計摘要"
    for r in dataframe_to_rows(df_s, index=False, header=True):
        ws_summary.append(r)
    pie = PieChart()
    pie.title = "各類別保養金額比例圖"
    data = Reference(ws_summary, min_col=2, min_row=1, max_row=len(df_s) + 1)
    labels = Reference(ws_summary, min_col=1, min_row=2, max_row=len(df_s) + 1)
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(labels)
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showVal = True
    ws_summary.add_chart(pie, "D2")
    
    ws_detail = wb.create_sheet(title="工單明細")
    for r in dataframe_to_rows(df_d, index=False, header=True):
        ws_detail.append(r)
    wb.save(output)
    return output.getvalue()

st.header("📋 系統歷史工單總表")
st.dataframe(df_detail, use_container_width=True)

excel_data = generate_excel(df_detail, df_detail.groupby("保養類別")["金額"].sum().reset_index())
st.download_button(
    label="📥 匯出並下載保養工單 Excel 檔",
    data=excel_data,
    file_name="汽車保養工單統計表.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
